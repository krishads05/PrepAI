"""Deep extraction from PDFs (incl. scanned/handwritten pages), PPTX (text, tables, notes, grouped shapes,
images), legacy PPT (via LibreOffice) and DOCX (paragraphs + tables). OCR/vision results are cached on disk.
MODE 'cache': use cached OCR only (fast, never blocks the app) and record what is pending.
MODE 'run'  : perform OCR/vision for anything not yet cached."""
import hashlib, json, os, pathlib, re, shutil, subprocess
from . import vision

CACHE = pathlib.Path(os.getenv("PREPAI_CACHE", ".prepai_cache"))
MIN_PAGE_TEXT = 40
MODE = "cache"
ENGINE = "fast"                                  # 'fast' (OCR + equation OCR) or 'vision' (local vision model)
INCLUDE_SCANNED = True                           # read scanned/handwritten PDF pages (papers: yes; notes: optional)
pending = []                                     # [(kind, label)] still to read

def _sha(b):
    return hashlib.sha1(b).hexdigest()

def _file_sha(p):
    return _sha(pathlib.Path(p).read_bytes())

def _get(key):
    f = CACHE / "ocr" / f"{key}.json"
    return json.loads(f.read_text(encoding="utf-8")) if f.exists() else None

def _put(key, text, engine, label):
    (CACHE / "ocr").mkdir(parents=True, exist_ok=True)
    (CACHE / "ocr" / f"{key}.json").write_text(json.dumps(dict(text=text, engine=engine, label=label)), encoding="utf-8")

def _ocr_cached(key, get_bytes, kind, label):
    c = _get(key)
    if c is not None:
        upgrade = (MODE == "run" and ENGINE == "vision" and not c["engine"].startswith("vision")
                   and vision.vision_model())
        if not upgrade:
            return c["text"]
    if kind == "page" and not INCLUDE_SCANNED:
        if MODE == "run":
            return ""                            # skipped on purpose (e.g. handwritten notes)
    if MODE != "run":
        pending.append((kind, label)); return ""
    text, engine = vision.ocr(get_bytes(), kind, engine=ENGINE)
    if engine != "none":
        _put(key, text, engine, label)
    return text

# ---------------------------------------------------------------- PDF
def _render_page(path, i, dpi=170):
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        with fitz.open(path) as d:
            return d[i].get_pixmap(dpi=dpi).tobytes("png")
    except ImportError:
        out = CACHE / "render"; out.mkdir(parents=True, exist_ok=True)
        stem = out / f"{_file_sha(path)[:10]}_{i}"
        subprocess.run(["pdftoppm", "-r", str(dpi), "-f", str(i + 1), "-l", str(i + 1), "-png", "-singlefile",
                        str(path), str(stem)], check=True, capture_output=True)
        return pathlib.Path(f"{stem}.png").read_bytes()

def _pdf_text_pages(path):
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        with fitz.open(path) as d:
            return [pg.get_text() for pg in d]
    except ImportError:
        import pdfplumber
        with pdfplumber.open(path) as d:
            return [pg.extract_text() or "" for pg in d.pages]

def pdf_pages(path):
    fsha = _file_sha(path)[:16]; name = pathlib.Path(path).name
    out = []
    for i, t in enumerate(_pdf_text_pages(path)):
        if len(t.strip()) < MIN_PAGE_TEXT:          # scanned / handwritten page
            t = _ocr_cached(f"pdf_{fsha}_{i}", lambda i=i: _render_page(path, i), "page", f"{name} p.{i + 1}")
        out.append((f"p. {i + 1}", t))
    return out

# ---------------------------------------------------------------- PPTX / PPT
def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

def _xml_lines(element):
    """Every DrawingML paragraph incl. Office Math equations, tables, grouped shapes; math italics -> plain."""
    import unicodedata
    lines = []
    for para in element.iter(A_NS + "p"):
        parts = [t.text for t in para.iter() if t.tag in (A_NS + "t", M_NS + "t") and t.text]
        line = unicodedata.normalize("NFKC", "".join(parts)).replace("\xa0", " ").strip()
        if line in ("<number>", "‹#›", "<#>") or re.fullmatch(r"\d{1,3}", line or ""):
            continue                                   # slide-number placeholders
        if line and (not lines or lines[-1] != line):
            lines.append(line)
    return lines

def pptx_pages(path, min_px=(120, 60), min_bytes=4000):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path); name = pathlib.Path(path).name
    seen_imgs, out = set(), []
    for n, slide in enumerate(prs.slides, 1):
        parts = _xml_lines(slide._element)
        for sh in _walk(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = sh.image
                except Exception:
                    continue
                w, h = img.size
                if img.sha1 in seen_imgs or w < min_px[0] or h < min_px[1] or len(img.blob) < min_bytes:
                    continue                          # logos repeated on every slide, icons, bullets
                seen_imgs.add(img.sha1)
                t = _ocr_cached(f"img_{img.sha1}", lambda b=img.blob: b, "image", f"{name} slide {n} image")
                if t.strip():
                    parts.append(f"[Image] {t.strip()}")
        if slide.has_notes_slide:
            notes = _xml_lines(slide.notes_slide._element)
            if notes:
                parts.append("[Notes] " + " ".join(notes))
        out.append((f"slide {n}", "\n".join(parts)))
    return out

def soffice_path():
    p = shutil.which("soffice") or shutil.which("libreoffice")
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if not p and os.path.exists(c):
            p = c
    return p

def convert_ppt(path):
    out = CACHE / "converted"; out.mkdir(parents=True, exist_ok=True)
    target = out / (pathlib.Path(path).stem + ".pptx")
    if not target.exists():
        exe = soffice_path()
        if not exe:
            raise RuntimeError("legacy .ppt needs LibreOffice (or open it in PowerPoint and Save As .pptx)")
        subprocess.run([exe, "--headless", "--convert-to", "pptx", "--outdir", str(out), str(path)],
                       check=True, capture_output=True, timeout=300)
    return target

# ---------------------------------------------------------------- DOCX
def docx_text(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        last = None
        for row in t.rows:
            for c in row.cells:
                x = c.text.strip()
                if x and x != last:                   # merged cells repeat their text
                    parts.append(x); last = x
    return "\n".join(parts)

# ---------------------------------------------------------------- entry points
def pages(path):
    p = pathlib.Path(path); ext = p.suffix.lower()
    if ext == ".pdf":
        return pdf_pages(p)
    if ext == ".pptx":
        return pptx_pages(p)
    if ext == ".ppt":
        return pptx_pages(convert_ppt(p))
    if ext == ".docx":
        t = docx_text(p)
    else:
        t = p.read_text(encoding="utf-8", errors="ignore")
    blocks = [b for b in re.split(r"\n\s*\n", t) if b.strip()]
    if len(blocks) <= 1:
        blocks = [b for b in t.splitlines() if b.strip()]
    return [(f"para {i + 1}", b) for i, b in enumerate(blocks)]

def text(path):
    p = pathlib.Path(path)
    if p.suffix.lower() == ".docx":
        return docx_text(p)
    return "\n\n".join(t for _, t in pages(p))

def process_folder(folder, progress=None, include_scanned=False, engine="fast"):
    """Read slide images (and optionally scanned/handwritten pages) in a folder. Resumable via the cache."""
    global MODE, ENGINE, INCLUDE_SCANNED
    todo = pending_report(folder)
    n = len(todo["images"]) + (len(todo["pages"]) if include_scanned else 0)
    files = sorted(x for x in pathlib.Path(folder).iterdir() if x.is_file() and not x.name.startswith("."))
    done = 0
    orig = vision.ocr
    def counted(img, kind="page", engine="auto", _o=orig):
        nonlocal done
        r = _o(img, kind, engine); done += 1
        if progress:
            progress(done, n, cur_name)
        return r
    MODE, ENGINE, INCLUDE_SCANNED = "run", engine, include_scanned
    vision.ocr = counted
    try:
        for f in files:
            cur_name = f.name
            try:
                pages(f)
            except Exception as e:
                print(f"[WARN] {f.name}: {e}")
    finally:
        vision.ocr = orig
        MODE, ENGINE, INCLUDE_SCANNED = "cache", "fast", True
    left = pending_report(folder)
    return dict(total=n, done=done, images_left=len(left["images"]), pages_left=len(left["pages"]))

def pending_report(folder):
    """{'images': [...], 'pages': [...]} still unread (fast: cache-only pass)."""
    pending.clear()
    for f in sorted(pathlib.Path(folder).iterdir()):
        if f.is_file() and not f.name.startswith("."):
            try:
                pages(f)
            except Exception:
                pass
    r = dict(images=[l for k, l in pending if k == "image"], pages=[l for k, l in pending if k == "page"])
    pending.clear()
    return r
